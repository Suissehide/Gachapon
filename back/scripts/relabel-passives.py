#!/usr/bin/env python3
"""Aligne la colonne Passif du classeur sur les libellés de passives.ts.

Le format de la colonne est `Libellé français (CLÉ)`. Seul le libellé change,
la clé fait foi et sert d'appariement.

Usage : python3 back/scripts/relabel-passives.py [--apply]
"""
import re
import shutil
import sys
from pathlib import Path

import openpyxl

XLSX = Path('/Volumes/Elements/Gachapon/tcg_kit.xlsx')
COPIE = XLSX.with_name('tcg_kit.relabeled.xlsx')

NOUVEAUX_LIBELLES = {
    'CRIT': 'Précision',
    'PIERCE': 'Perce-armure',
    'VAMPIRISM': 'Vampirisme',
    'VIGOR': 'Second souffle',
    'HASTE': 'Célérité',
    'FORTIFY': 'Fortification',
    'EMPOWER': 'Puissance',
}

apply = '--apply' in sys.argv
# La copie n'est écrite sur disque qu'en mode --apply : un essai à blanc
# charge l'original en lecture seule (jamais réécrit) pour que le message
# « copie non sauvegardée » soit vrai (même correctif que rebalance-cards.py).
if apply:
    shutil.copy(XLSX, COPIE)
    wb = openpyxl.load_workbook(COPIE, data_only=False)
else:
    wb = openpyxl.load_workbook(XLSX, data_only=False)
ws = wb['Production']
entetes = [c.value for c in ws[3]]
col_passif = entetes.index('Passif') + 1

modifiees = 0
for ligne in range(4, ws.max_row + 1):
    cellule = ws.cell(row=ligne, column=col_passif)
    valeur = cellule.value
    if not valeur:
        continue
    correspondance = re.match(r'^(.*)\s*\(([A-Z_]+)\)$', str(valeur).strip())
    if not correspondance:
        print(f'  ligne {ligne} : format inattendu -> {valeur!r}')
        continue
    cle = correspondance.group(2)
    if cle in NOUVEAUX_LIBELLES:
        nouveau = f'{NOUVEAUX_LIBELLES[cle]} ({cle})'
        if nouveau != valeur:
            cellule.value = nouveau
            modifiees += 1

print(f'{modifiees} libellés mis à jour -> {COPIE}')
if apply:
    wb.save(COPIE)
else:
    print('  (essai à blanc, copie non sauvegardée)')
