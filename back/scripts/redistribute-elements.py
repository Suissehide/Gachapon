#!/usr/bin/env python3
"""Ramène LIGHT / DARK à ~10 % du roster.

Pourquoi : un tiers des cartes étaient LIGHT ou DARK, neutres face au cycle
FIRE ▶ NATURE ▶ EARTH ▶ WATER. La paire devient rare (et, à terme, réservée à
la capsule astrale) ; les autres basculent dans le cycle.

Choix des cartes, sur le nom + les tags booru du classeur :
  - on garde autant de LIGHT que de DARK, réparties sur toutes les familles
    et toutes les raretés (quota proportionnel au poids de la rareté parmi
    LIGHT/DARK) ; dans ces quotas, les cartes au thème lumineux / ténébreux
    le plus marqué, même quand le thème ne colle pas toujours ;
  - les autres prennent l'élément du cycle que leur thème évoque, tant que
    cet élément n'a pas atteint sa part du roster ; sinon le suivant dans
    l'ordre de leur thème, à égalité celui que leur famille possède le moins
    (le cycle reste équilibré).

Mêmes quatre sources que rebalance-cards.py : classeur (sauvegardé, la
sauvegarde sert de verrou), cards-data.json, Humains de cards.definitions.ts,
migration de données reconnaissant les cartes par leur image.

Usage : python3 back/scripts/redistribute-elements.py [--apply]
Sans --apply, affiche le tableau de ce qui changerait.
"""
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

XLSX = Path('/Volumes/Elements/Gachapon/tcg_kit.xlsx')
SAUVEGARDE = XLSX.with_name('tcg_kit.avant-elements-2026-10-30.xlsx')
SCRIPTS = Path(__file__).resolve().parent
JSON_PATH = SCRIPTS / 'import-cards' / 'cards-data.json'
HUMAINS_TS = SCRIPTS.parent / 'src/main/domain/content/cards.definitions.ts'
MIGRATION = (
    SCRIPTS.parent
    / 'prisma/migrations/20261030120000_elements_lumiere_tenebres_rares/migration.sql'
)

PART_CIBLE = 0.10
CYCLE = ('FIRE', 'WATER', 'NATURE', 'EARTH')
PAIRE = ('LIGHT', 'DARK')
LIBELLE = {
    'FIRE': 'Feu (FIRE)',
    'WATER': 'Eau (WATER)',
    'NATURE': 'Nature (NATURE)',
    'EARTH': 'Terre (EARTH)',
    'LIGHT': 'Lumière (LIGHT)',
    'DARK': 'Ténèbres (DARK)',
}
RARETES = ('Commun', 'Uncommun', 'Rare', 'Épique', 'Légendaire')

MOTS = {
    'FIRE': r'fire|flame|ember|lava|magma|burn|torch|inferno|phoenix|crimson|'
    r'red (?:scales|hair|eyes)|feu|flamm|braise|cendr|pyr|ign|brasier|ardent',
    'WATER': r'water|ice|frost|snow|sea|ocean|wave|aqua|rain|glacier|coral|'
    r'blue (?:scales|hair|eyes)|siren|mermaid|glac|givr|neige|onde|mar[ée]|'
    r'flot|ondin|pluie|abysse marine',
    'NATURE': r'leaf|leaves|flower|vine|forest|moss|wood|thorn|druid|hunter|'
    r'fur|petal|antler|green (?:scales|hair|eyes)|feuill|fleur|sylv|ronc|'
    r'foug|mousse|bois|liane|épine|sève|bruy',
    'EARTH': r'stone|rock|earth|crystal|gem|sand|mountain|iron|bronze|metal|'
    r'heavy armor|golem|miner|brown (?:scales|hair)|pierre|roc|terre|sable|'
    r'granit|bronz|\bfer\b|cristal|mont',
    'LIGHT': r'holy|halo|radiant|white|glow|sun|golden|angel|priest|paladin|'
    r'divine|star|feather|lumi|aube|solaire|aur[eéo]|clart|céleste|'
    r'\bsol\b|séraph|astre',
    'DARK': r'shadow|dark|black|skull|demon|void|necro|night|moon|horn|blood|'
    r'vampire|undead|bone|curse|ombre|nuit|noir|sombre|néant|mort|lune|'
    r'sang|maudi|spectr|banshee|wraith|ghost|lich|fallen|déchu|fantôme',
}
# Familles dont l'identité EST la paire : leur thème pèse en plus des tags.
FAMILLE = {
    'Anges': ('LIGHT', 3),
    'Démons': ('DARK', 3),
    'Morts-vivants': ('DARK', 3),
}

apply = '--apply' in sys.argv
if apply and SAUVEGARDE.exists():
    raise SystemExit(f'✗ {SAUVEGARDE.name} existe : la passe a déjà eu lieu.')


def cle(valeur: str) -> str:
    return re.search(r'\(([A-Z]+)\)\s*$', valeur).group(1)


def scores(nom: str, tags: str, famille: str) -> dict[str, int]:
    texte = f'{nom} {tags}'.lower().replace('light armor', '')
    s = {el: len(re.findall(m, texte)) for el, m in MOTS.items()}
    if famille in FAMILLE:
        el, bonus = FAMILLE[famille]
        s[el] += bonus
    return s


# --- classeur -------------------------------------------------------------
wb = openpyxl.load_workbook(XLSX, data_only=False)
ws = wb['Production']
entetes = [c.value for c in ws[3]]
col = {nom: i + 1 for i, nom in enumerate(entetes)}


def val(ligne: int, nom: str):
    return ws.cell(row=ligne, column=col[nom]).value


cartes = []
for ligne in range(4, ws.max_row + 1):
    ident = val(ligne, 'ID')
    if not ident or val(ligne, 'Famille') == 'Boss' or not val(ligne, 'Élément'):
        continue
    cartes.append({
        'ligne': ligne,
        'id': ident,
        'nom': val(ligne, 'Nom de la carte'),
        'famille': val(ligne, 'Famille'),
        'rarete': val(ligne, 'Rareté'),
        'avant': cle(val(ligne, 'Élément')),
        's': scores(val(ligne, 'Nom de la carte'), val(ligne, 'Sujet (tags booru)') or '',
                    val(ligne, 'Famille')),
    })

paire = [c for c in cartes if c['avant'] in PAIRE]
cible = round(len(cartes) * PART_CIBLE)
familles = sorted({c['famille'] for c in cartes})
par_rarete = Counter(c['rarete'] for c in paire)

# Quota par rareté, proportionnel au poids de la rareté parmi LIGHT/DARK
# (le reste de l'arrondi va aux communes).
quota = {r: round(par_rarete[r] * cible / len(paire)) for r in RARETES[1:]}
quota['Commun'] = cible - sum(quota.values())

# Quota par famille et par élément : toutes les familles en ont, autant de
# LIGHT que de DARK. Les familles les plus fournies en paire prennent le
# reste de la division ; une famille impaire penche vers son thème, en
# alternant pour que les deux totaux restent égaux.
base, reste = divmod(cible, len(familles))
taille = Counter(c['famille'] for c in paire)
ordre = sorted(familles, key=lambda f: -taille[f])
part = {f: base + (i < reste) for i, f in enumerate(ordre)}


def penchant(f: str) -> int:
    return sum(c['s']['LIGHT'] - c['s']['DARK'] for c in paire if c['famille'] == f)


cap = {}
impaires = sorted((f for f in familles if part[f] % 2), key=penchant, reverse=True)
for f in familles:
    cap[f] = Counter({'LIGHT': part[f] // 2, 'DARK': part[f] // 2})
for i, f in enumerate(impaires):
    cap[f][PAIRE[0] if i < len(impaires) // 2 else PAIRE[1]] += 1


def marque(c, el):
    """Thème de l'élément visé face au meilleur thème du cycle."""
    return c['s'][el] - max(c['s'][x] for x in CYCLE)


# Des raretés les plus rares aux communes : à chaque place, la carte au thème
# le plus marqué parmi celles dont la famille a encore de la place pour cet
# élément. Les deux éléments alternent pour rester à égalité.
nouveau: dict[str, str] = {}
for rarete in reversed(RARETES):
    for i in range(quota[rarete]):
        el = PAIRE[i % 2] if rarete != 'Commun' or i % 2 == 0 else PAIRE[1]
        choix = [c for c in paire if c['rarete'] == rarete
                 and c['id'] not in nouveau and cap[c['famille']][el] > 0]
        if not choix:
            el = PAIRE[1 - PAIRE.index(el)]
            choix = [c for c in paire if c['rarete'] == rarete
                     and c['id'] not in nouveau and cap[c['famille']][el] > 0]
        if not choix:
            continue
        meilleur = max(choix, key=lambda c: (marque(c, el), c['avant'] == el))
        nouveau[meilleur['id']] = el
        cap[meilleur['famille']][el] -= 1

manque = sum(sum(v.values()) for v in cap.values())
if manque:
    # Places restées vides (une famille sans carte de la rareté voulue) :
    # on les comble toutes raretés confondues.
    for f in familles:
        for el in PAIRE:
            while cap[f][el] > 0:
                choix = [c for c in paire if c['famille'] == f and c['id'] not in nouveau]
                if not choix:
                    raise SystemExit(f'✗ {f} : pas assez de cartes LIGHT/DARK')
                meilleur = max(choix, key=lambda c: (marque(c, el), c['avant'] == el))
                nouveau[meilleur['id']] = el
                cap[f][el] -= 1

# Le reste bascule dans le cycle. Les thèmes les plus marqués choisissent en
# premier ; un élément plein (sa part du roster atteinte) ne prend plus personne.
compte = defaultdict(Counter)  # famille -> élément
total = Counter()
for c in cartes:
    if c['avant'] in CYCLE:
        compte[c['famille']][c['avant']] += 1
        total[c['avant']] += 1
plafond = -(-(len(cartes) - cible) // len(CYCLE))
for c in sorted(paire, key=lambda c: -max(c['s'][x] for x in CYCLE)):
    if c['id'] in nouveau:
        continue
    libres = [x for x in CYCLE if total[x] < plafond]
    el = max(libres, key=lambda x: (c['s'][x], -compte[c['famille']][x], -total[x]))
    nouveau[c['id']] = el
    compte[c['famille']][el] += 1
    total[el] += 1

changes = {i: e for i, e in nouveau.items()
           if e != next(c['avant'] for c in cartes if c['id'] == i)}

# --- rapport --------------------------------------------------------------
apres = Counter(nouveau.get(c['id'], c['avant']) for c in cartes)
avant = Counter(c['avant'] for c in cartes)
print(f'{len(cartes)} cartes, cible LIGHT+DARK = {cible}')
for el in CYCLE + PAIRE:
    print(f'  {el:<7} {avant[el]:>4} -> {apres[el]:>4}')
print('\nLIGHT/DARK gardées, par rareté :')
for rarete in RARETES:
    k = Counter(nouveau[c['id']] for c in paire
                if c['rarete'] == rarete and nouveau[c['id']] in PAIRE)
    print(f'  {rarete:<11} LIGHT {k["LIGHT"]:>2}  DARK {k["DARK"]:>2}')
print('\nID        rareté      famille        avant  -> après   nom')
for c in sorted(paire, key=lambda c: (nouveau[c['id']] not in PAIRE,
                                      RARETES.index(c['rarete']), c['id'])):
    print(f"{c['id']:<9} {c['rarete']:<11} {c['famille']:<14} "
          f"{c['avant']:<6} -> {nouveau[c['id']]:<7}  {c['nom']}")

if not apply:
    print('\n(essai à blanc, rien écrit)')
    sys.exit()

for c in cartes:
    if c['id'] in changes:
        ws.cell(row=c['ligne'], column=col['Élément']).value = LIBELLE[changes[c['id']]]
shutil.copy(XLSX, SAUVEGARDE)
wb.save(XLSX)
print(f'classeur : {len(changes)} cartes, sauvegarde -> {SAUVEGARDE.name}')

donnees = json.loads(JSON_PATH.read_text())
n = 0
for lot in donnees.values():
    for c in lot:
        if c['id'] in changes:
            c['element'] = changes[c['id']]
            n += 1
JSON_PATH.write_text(json.dumps(donnees, ensure_ascii=False, indent=2) + '\n')
print(f'cards-data.json : {n} cartes')

source = HUMAINS_TS.read_text()
BLOC = re.compile(r"(id: '(HUM-\d{3})',.*?element: ')[A-Z]+(')", re.S)
source, h = BLOC.subn(
    lambda m: f'{m.group(1)}{changes.get(m.group(2), m.group(0).rsplit(chr(39), 2)[1])}{m.group(3)}',
    source,
)
HUMAINS_TS.write_text(source)
print(f'cards.definitions.ts : {sum(i.startswith("HUM-") for i in changes)} Humains')

lignes = ',\n'.join(f"  ('{i}', '{e}')" for i, e in sorted(changes.items()))
MIGRATION.parent.mkdir(parents=True, exist_ok=True)
MIGRATION.write_text(f"""-- LIGHT / DARK ramenées à ~10 % du roster (scripts/redistribute-elements.py).
--
-- Un tiers des cartes portaient la paire, neutre face au cycle : elle
-- n'apportait rien. Il en reste 30 LIGHT et 30 DARK, réparties sur toutes les
-- familles et toutes les raretés ; les autres prennent l'élément du cycle que
-- leur thème évoque, le cycle restant équilibré.
--
-- Migration de DONNÉES : le déploiement ne seede jamais. Les cartes sont
-- reconnues par leur image (`<dossier>/<ID>.png`) ; une carte pas encore
-- importée est ignorée, l'import la créera avec cards-data.json. Idempotent.

UPDATE "Card" c
SET "element" = v.element::"CardElement"
FROM (VALUES
{lignes}
) AS v(id, element)
WHERE c."imageUrl" LIKE '%/' || v.id || '.png';
""")
print(f'migration -> {MIGRATION.relative_to(SCRIPTS.parent)}')
